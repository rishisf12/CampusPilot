# -*- coding: utf-8 -*-
"""Generate a Japan-targeted, English-only resume (DOCX) with a new layout."""
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

NAVY = RGBColor(0x1F, 0x3A, 0x5F)
SLATE = RGBColor(0x44, 0x54, 0x6A)
GRAY = RGBColor(0x6B, 0x72, 0x80)
BLACK = RGBColor(0x21, 0x25, 0x29)

doc = Document()

# ---------- page setup: A4 ----------
sec = doc.sections[0]
sec.page_width = Cm(21.0)
sec.page_height = Cm(29.7)
sec.top_margin = Cm(1.6)
sec.bottom_margin = Cm(1.4)
sec.left_margin = Cm(1.9)
sec.right_margin = Cm(1.9)

style = doc.styles['Normal']
style.font.name = 'Calibri'
style.font.size = Pt(10)
style.font.color.rgb = BLACK
pf = style.paragraph_format
pf.space_after = Pt(3)
pf.space_before = Pt(0)
pf.line_spacing = 1.04

def set_cell_bg(cell, hex_color):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:fill'), hex_color)
    tcPr.append(shd)

def para_border_bottom(paragraph, color='1F3A5F', size='6', space='4'):
    pPr = paragraph._p.get_or_add_pPr()
    pBdr = OxmlElement('w:pBdr')
    bottom = OxmlElement('w:bottom')
    bottom.set(qn('w:val'), 'single')
    bottom.set(qn('w:sz'), size)
    bottom.set(qn('w:space'), space)
    bottom.set(qn('w:color'), color)
    pBdr.append(bottom)
    pPr.append(pBdr)

def add_run(p, text, size=10, bold=False, italic=False, color=BLACK, font='Calibri'):
    r = p.add_run(text)
    r.font.name = font
    r.font.size = Pt(size)
    r.bold = bold
    r.italic = italic
    r.font.color.rgb = color
    return r

def heading(text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(9)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.keep_with_next = True
    add_run(p, text.upper(), size=10.5, bold=True, color=NAVY)
    para_border_bottom(p)
    return p

def body(text_runs, space_after=3, align=None):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(space_after)
    if align:
        p.alignment = align
    for text, kw in text_runs:
        add_run(p, text, **kw)
    return p

def bullet(text_runs, indent=0.32, hang=0.32, space_after=2):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Cm(indent)
    p.paragraph_format.first_line_indent = Cm(-hang)
    p.paragraph_format.space_after = Pt(space_after)
    add_run(p, '\u25aa  ', size=9, color=NAVY, bold=True)
    for text, kw in text_runs:
        add_run(p, text, **kw)
    return p

# ================= HEADER =================
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.paragraph_format.space_after = Pt(1)
add_run(p, 'RISHIKESH SAROJ', size=20, bold=True, color=NAVY)

p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.paragraph_format.space_after = Pt(2)
add_run(p, 'B.Tech Mechanical Engineering · Software & AI/ML Engineer', size=10.5, bold=True, color=SLATE)

p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.paragraph_format.space_after = Pt(4)
add_run(p, '+91 8299885894   |   rishisf120@gmail.com   |   github.com/rishisf12   |   linkedin.com/in/rishisf12', size=9, color=GRAY)
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.paragraph_format.space_after = Pt(6)
add_run(p, 'PDPM IIITDM Jabalpur, India  ·  Expected graduation 2028  ·  Open to relocation — seeking internship / full-time roles in Japan', size=9, italic=True, color=GRAY)

div = doc.add_paragraph()
div.paragraph_format.space_after = Pt(2)
para_border_bottom(div, color='1F3A5F', size='8', space='1')

# ================= PROFESSIONAL SUMMARY =================
heading('Professional Summary')
body([
    ('Mechanical engineering undergraduate building end-to-end ML and software applications '
     '(FastAPI, PostgreSQL, React, Docker, Scikit-learn, LLMs/OCR). CampusPilot serves ~1,500 students '
     'and is expanding to mobile via React Native. Strong full-stack fundamentals with 60+ automated tests '
     'and containerized deployments.', {'size': 9.5}),
    ('Seeking a software engineering / AI-ML internship or full-time role in Japan. Eager to contribute '
     'to a product-driven engineering team and grow within a Japanese company long-term.', {'size': 9.5, 'bold': True}),
], space_after=6)

# ================= TECHNICAL SKILLS (two-column grid) =================
heading('Technical Skills')
skills = [
    ('Languages', 'Python (primary) · JavaScript/TypeScript · SQL · C++'),
    ('AI / ML', 'Scikit-learn · OpenCV · Tesseract OCR · Hugging Face (Phi-3, Mistral, Gemma) · Pandas · NumPy · Matplotlib'),
    ('Backend', 'FastAPI · REST APIs · SQLModel · PostgreSQL · Alembic · Docker · pytest'),
    ('Frontend / Tools', 'React · Vite · Tailwind CSS · Recharts · Git/GitHub'),
]
tbl = doc.add_table(rows=2, cols=2)
tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
tbl.autofit = False
tbl.style = doc.styles['Normal Table']
widths = [Cm(8.6), Cm(8.6)]
for r_i, row in enumerate(tbl.rows):
    for c_i, cell in enumerate(row.cells):
        cell.width = widths[c_i]
        cell.paragraphs[0].paragraph_format.space_after = Pt(1)
        cell.paragraphs[0].paragraph_format.space_before = Pt(0)
        label, value = skills[r_i * 2 + c_i]
        p = cell.paragraphs[0]
        add_run(p, label + '  ', size=8.5, bold=True, color=NAVY)
        add_run(p, value, size=8.5, color=SLATE)
        set_cell_bg(cell, 'F2F5F9')
for row in tbl.rows:
    for cell in row.cells:
        tcPr = cell._tc.get_or_add_tcPr()
        tcMar = OxmlElement('w:tcMar')
        for side, w in (('top', '20'), ('bottom', '20'), ('left', '80'), ('right', '80')):
            el = OxmlElement(f'w:{side}')
            el.set(qn('w:w'), w)
            el.set(qn('w:type'), 'dxa')
            tcMar.append(el)
        tcPr.append(tcMar)
tblPr = tbl._tbl.tblPr
borders = OxmlElement('w:tblBorders')
for edge in ('top', 'left', 'bottom', 'right', 'insideH', 'insideV'):
    e = OxmlElement(f'w:{edge}')
    e.set(qn('w:val'), 'none')
    borders.append(e)
tblPr.append(borders)

doc.add_paragraph().paragraph_format.space_after = Pt(0)

# ================= PROJECTS =================
heading('Projects')

def project(title, meta, tech, bullets):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(5)
    p.paragraph_format.space_after = Pt(1)
    p.paragraph_format.keep_with_next = True
    add_run(p, title, size=10.5, bold=True, color=BLACK)
    p2 = doc.add_paragraph()
    p2.paragraph_format.space_after = Pt(1)
    add_run(p2, meta, size=8.5, italic=True, color=GRAY)
    p3 = doc.add_paragraph()
    p3.paragraph_format.space_after = Pt(1)
    add_run(p3, tech, size=8.5, color=SLATE)
    for b in bullets:
        bullet(b, space_after=1)

project(
    'CampusPilot — Academic management platform (~1,500 students)',
    'In progress  ·  github.com/rishisf12/CampusPilot',
    'Python · FastAPI · Hugging Face (LLMs) · Tesseract OCR · ReportLab · React',
    [
        [('Built an OCR + LLM pipeline that converts exam seating PDFs into JSON for roll-number lookup and generates personalized PDF timetables.', {'size': 9})],
        [('Built attendance tracking with Safe/Warning/Critical flags, a “classes you can miss” calculator, and live current/next class views from PDF/CSV timetables.', {'size': 9})],
        [('Designed a layered architecture (routes, services, models, OCR modules) backed by 63+ unit tests.', {'size': 9})],
    ]
)

project(
    'pocketF — Student budgeting app (solo project)',
    'In progress  ·  github.com/rishisf12',
    'Python · FastAPI · PostgreSQL · Pandas · Matplotlib · Docker Compose · pytest · React',
    [
        [('Built a CSV importer using Decimal for exact money math, flagging bad rows across 9 error types.', {'size': 9})],
        [('Built burn-rate analytics and a “money runs out on” forecast with SAFE/ALERT verdicts and auto-generated charts.', {'size': 9})],
        [('Containerized API + PostgreSQL with Docker Compose; tested forecast edge cases with pytest.', {'size': 9})],
    ]
)

project(
    'UrTeam — Hackathon matchmaking platform',
    'github.com/rishisf12',
    'Python · FastAPI · SQLModel · PostgreSQL/SQLite · React · Vite · Tailwind CSS',
    [
        [('Matched students into hackathon teams via skill overlap using set-based Jaccard similarity (no ML dependencies).', {'size': 9})],
        [('Designed 15+ REST endpoints (student CRUD with 409 duplicate detection, ?tech= idea filter), domain-validated emails, SQLite/PostgreSQL dual-database with auto table creation.', {'size': 9})],
        [('Built a chat approval flow with owner-verified access control; 8 reusable React components with inline Tailwind CSS.', {'size': 9})],
    ]
)

project(
    'Master-DSA — DSA practice & language comparison repo',
    'github.com/rishisf12/master-dsa',
    'C++ · Python · Git/GitHub',
    [
        [('Practicing DSA problems and tracking learning progress across topics.', {'size': 9})],
        [('Comparing C++ and Python code execution, performance, and limits side by side.', {'size': 9})],
        [('Maintaining a personal DSA learning log with implemented solutions and notes.', {'size': 9})],
    ]
)

# ================= EDUCATION =================
heading('Education')
body([
    ('PDPM IIITDM Jabalpur, Madhya Pradesh, India', {'size': 9.5, 'bold': True}),
    ('   B.Tech, Mechanical Engineering   ·   Expected 2028   ·   CPI: 6.9 / 10', {'size': 9, 'color': SLATE}),
    ('Class XII: 80%   |   Class X: 92%', {'size': 9, 'color': GRAY}),
], space_after=6)

# ================= ACHIEVEMENTS =================
heading('Achievements')
bullet([('Smart India Hackathon 2026 — ', {'size': 9, 'bold': True}), ('Team of 6; proposed SkyGuard AI (SIH26073), Isolation Forest + SHAP anomaly detection to separate genuine weather events from sensor faults in Automatic Weather Stations.', {'size': 9})])
bullet([('Patent — ', {'size': 9, 'bold': True}), ('Co-inventor, Indian Patent Application No. 202621096766 A: adjustable hybrid welding device with integrated safety measures.', {'size': 9})])
bullet([('Jagrati (Volunteer) — ', {'size': 9, 'bold': True}), ('Teach underprivileged village children weekly; help plan the student-run initiative’s activities.', {'size': 9})])

note = doc.add_paragraph()
note.paragraph_format.space_before = Pt(10)
note.paragraph_format.space_after = Pt(0)
para_border_bottom(note, color='C9D2DE', size='4', space='2')
add_run(note, 'References available upon request.  ·  Portfolio & code: github.com/rishisf12', size=8, italic=True, color=GRAY)

out = r'C:\Users\Appex\Documents\Default Project\CampusPilot\Rishikesh_Resume_Japan.docx'
doc.save(out)
print('Saved:', out)
