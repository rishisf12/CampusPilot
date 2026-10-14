"""Dump Claude's Excel to a normalized JSON for comparison."""
import json
import openpyxl
from pathlib import Path

excel_path = Path(r"F:\UserFolders\Downloads\Timetable_Aug-Nov_2026_by_Semester_Branch.xlsx")
wb = openpyxl.load_workbook(excel_path)

# Collect All Entries sheet (the normalized flat table)
ws = wb["All Entries"]
headers = [cell.value for cell in ws[1]]
all_entries = []
for row in ws.iter_rows(min_row=2, values_only=True):
    if not any(row):
        continue
    all_entries.append(dict(zip(headers, row)))

# Also collect each semester-branch sheet
sheets_data = {}
for sheet_name in wb.sheetnames:
    if sheet_name in ("Index", "All Entries", "Check Items"):
        continue
    ws = wb[sheet_name]
    rows = []
    for row in ws.iter_rows(min_row=1, values_only=True):
        rows.append([cell for cell in row])
    sheets_data[sheet_name] = rows

output = {
    "all_entries": all_entries,
    "sheets_data": sheets_data,
}

out_path = Path(r"F:\UserFolders\Downloads\claude_excel_dump.json")
out_path.write_text(json.dumps(output, indent=2, default=str), encoding="utf-8")

print(f"All Entries rows: {len(all_entries)}")
print(f"Sheets: {len(sheets_data)}")
print(f"First 5 entries:")
for e in all_entries[:5]:
    print(" ", e)
print(f"Saved to {out_path}")
