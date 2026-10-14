"""Excel export for timetable data - generates multi-sheet workbook matching Claude's format."""
import logging
from datetime import time
from pathlib import Path
from typing import List, Dict, Any, Optional
from collections import OrderedDict

import openpyxl
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.utils import get_column_letter

from app.models import TimetableSlot
from app.core.config import DAYS_ORDER

logger = logging.getLogger(__name__)

# Semester display mapping
SEMESTER_DISPLAY = {
    1: "I",
    2: "II",
    3: "III",
    4: "IV",
    5: "V",
    6: "VI",
    7: "VII",
    8: "VIII",
}

# Branch display mapping for sheet names
BRANCH_SHEET_NAME = {
    "CSE A": "CSE-A",
    "CSE B": "CSE-B",
    "CSE-A": "CSE-A",
    "CSE-B": "CSE-B",
    "ECE": "ECE",
    "ME": "ME",
    "SM": "SM",
    "DS-E1": "DS-E1",
    "DS-E2": "DS-E2",
    "DS": "DS",
    "M.Tech (AI)": "MTech-AI",
    "M.Tech (DS)": "MTech-DS",
    "Ph.D. (CSE)": "PhD-CSE",
    "M.Tech (P&C)": "MTech-PC",
    "M.Tech (Nano)": "MTech-Nano",
    "M.Tech (Design)": "MTech-Design",
    "M.Tech (Mechtronics)": "MTech-Mechtronics",
    "Ph.D. (ME)": "PhD-ME",
    "MDes (Sem I)": "MDes-Sem-I",
    "MDes (Sem III)": "MDes-Sem-III",
    "Ph.D. (NS)": "PhD-NS",
}


def semester_to_roman(semester: int) -> str:
    """Convert semester number to Roman numeral for display."""
    return SEMESTER_DISPLAY.get(semester, str(semester))


def branch_to_sheet_name(branch: str) -> str:
    """Convert branch to sheet-safe name."""
    return BRANCH_SHEET_NAME.get(branch, branch.replace(" ", "-").replace("(", "").replace(")", "").replace("&", ""))


def get_sheet_name(semester: int, branch: str) -> str:
    """Generate sheet name like 'I-CSE-A', 'III-ECE', 'V-CSE-A', 'VII-CSE', 'PG-MTech-AI'."""
    sem_roman = semester_to_roman(semester)
    branch_short = branch_to_sheet_name(branch)
    
    if semester >= 9 or "M.Tech" in branch or "Ph.D." in branch or "MDes" in branch:
        return f"PG-{branch_short}"
    return f"{sem_roman}-{branch_short}"


def time_to_str(t: Optional[time]) -> str:
    """Convert time object to HH:MM-HH:MM string."""
    if not t:
        return ""
    # Assuming 1-hour slots
    start = t.strftime("%H:%M")
    end = (t.hour + 1, t.minute)
    end_str = f"{end[0]:02d}:{end[1]:02d}"
    return f"{start}-{end_str}"


def get_day_order_index(day: str) -> int:
    """Get sort index for day."""
    day_map = {"Mon": 0, "Tue": 1, "Wed": 2, "Thu": 3, "Fri": 4, "Sat": 5, "Sun": 6}
    return day_map.get(day[:3], 99)


class TimetableExcelExporter:
    """Exports timetable slots to multi-sheet Excel workbook."""
    
    def __init__(self, slots: List[TimetableSlot]):
        self.slots = slots
        self.workbook = openpyxl.Workbook()
        self._setup_styles()
    
    def _setup_styles(self):
        """Define cell styles."""
        self.header_font = Font(bold=True, size=11)
        self.title_font = Font(bold=True, size=14)
        self.subtitle_font = Font(bold=True, size=11, italic=True)
        self.day_font = Font(bold=True, size=11, color="FFFFFF")
        self.day_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
        self.header_fill = PatternFill(start_color="D9E2F3", end_color="D9E2F3", fill_type="solid")
        self.thin_border = Border(
            left=Side(style="thin"),
            right=Side(style="thin"),
            top=Side(style="thin"),
            bottom=Side(style="thin"),
        )
        self.wrap_alignment = Alignment(wrap_text=True, vertical="top")
        self.center_alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    
    def _apply_header_style(self, ws, row: int, max_col: int):
        """Apply header style to a row."""
        for col in range(1, max_col + 1):
            cell = ws.cell(row=row, column=col)
            cell.font = self.header_font
            cell.fill = self.header_fill
            cell.border = self.thin_border
            cell.alignment = self.center_alignment
    
    def _apply_day_header_style(self, ws, row: int, max_col: int):
        """Apply day section header style."""
        for col in range(1, max_col + 1):
            cell = ws.cell(row=row, column=col)
            cell.font = self.day_font
            cell.fill = self.day_fill
            cell.border = self.thin_border
            cell.alignment = self.center_alignment
        # Merge cells for day header
        ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=max_col)
    
    def _write_sheet_header(self, ws, semester: int, branch: str, row: int = 1) -> int:
        """Write sheet title and column headers. Returns next row index."""
        # Title
        sem_display = semester_to_roman(semester) if semester < 9 else "PG"
        ws.cell(row=row, column=1, value=f"Semester {sem_display} - {branch}").font = self.title_font
        row += 1
        ws.cell(row=row, column=1, value="Time Table Aug-Nov 2026 (Version III, effective 17 Aug 2026, updated 14 Aug 2026)").font = self.subtitle_font
        row += 2
        
        # Column headers
        headers = ["Course Code", "Instructor", "Room", "Time", "Branch", "Remarks"]
        for col, header in enumerate(headers, 1):
            ws.cell(row=row, column=col, value=header)
        self._apply_header_style(ws, row, len(headers))
        return row + 1
    
    def _write_day_section(self, ws, day: str, day_slots: List[Dict], start_row: int, branch: str) -> int:
        """Write a day's slots. Returns next row index."""
        if not day_slots:
            return start_row
        
        # Day header
        ws.cell(row=start_row, column=1, value=day.upper())
        self._apply_day_header_style(ws, start_row, 6)
        row = start_row + 1
        
        # Write slots
        for slot in day_slots:
            ws.cell(row=row, column=1, value=slot["course_code"])
            ws.cell(row=row, column=2, value=slot["instructor"])
            ws.cell(row=row, column=3, value=slot["room"])
            ws.cell(row=row, column=4, value=slot["time"])
            ws.cell(row=row, column=5, value=slot["branch"])
            ws.cell(row=row, column=6, value=slot["remarks"])
            
            for col in range(1, 7):
                cell = ws.cell(row=row, column=col)
                cell.border = self.thin_border
                cell.alignment = self.wrap_alignment
            row += 1
        
        # Blank row after day
        row += 1
        return row
    
    def build_semester_branch_sheet(self, semester: int, branch: str) -> openpyxl.worksheet.worksheet.Worksheet:
        """Build a single semester-branch sheet."""
        sheet_name = get_sheet_name(semester, branch)
        ws = self.workbook.create_sheet(title=sheet_name)
        
        # Filter slots for this semester and branch
        branch_slots = [s for s in self.slots if s.semester == semester and s.branch_or_program == branch]
        
        # Group by day
        day_slots: Dict[str, List[Dict]] = OrderedDict()
        for day in DAYS_ORDER:
            day_slots[day] = []
        
        for slot in branch_slots:
            day_key = slot.day[:3] if slot.day else "Mon"
            if day_key in day_slots:
                day_slots[day_key].append({
                    "course_code": slot.course_code,
                    "instructor": slot.instructor or "",
                    "room": slot.room or "",
                    "time": time_to_str(slot.start_time),
                    "branch": slot.branch_or_program,
                    "remarks": slot.remarks or "",
                })
        
        # Write sheet
        row = self._write_sheet_header(ws, semester, branch)
        
        for day in DAYS_ORDER:
            if day_slots[day]:
                row = self._write_day_section(ws, day, day_slots[day], row, branch)
        
        # Add hyperlink to Index
        ws.cell(row=1, column=8, value="<< Index").font = Font(italic=True, color="0000FF")
        ws.cell(row=1, column=8).hyperlink = "#'Index'!A1"
        
        # Adjust column widths
        ws.column_dimensions["A"].width = 18
        ws.column_dimensions["B"].width = 18
        ws.column_dimensions["C"].width = 15
        ws.column_dimensions["D"].width = 15
        ws.column_dimensions["E"].width = 15
        ws.column_dimensions["F"].width = 40
        ws.column_dimensions["H"].width = 12
        
        return ws
    
    def build_index_sheet(self) -> openpyxl.worksheet.worksheet.Worksheet:
        """Build the Index sheet."""
        ws = self.workbook.active
        ws.title = "Index"
        
        # Title
        ws.cell(row=1, column=1, value="Time Table Aug-Nov 2026 - rebuilt by semester / programme / branch").font = self.title_font
        ws.cell(row=2, column=1, value="Time Table Aug-Nov 2026 (Version III, effective 17 Aug 2026, updated 14 Aug 2026). Source: DOC-20260817-WA0002-print.PDF").font = self.subtitle_font
        ws.cell(row=3, column=1, value="Each sheet is organised Monday-Saturday. Columns: Course Code, Instructor, Room, Time, Branch (+ Remarks).").font = Font(italic=True, size=10)
        
        row = 5
        headers = ["Sheet (click to open)", "Level", "Branch", "No. of class entries"]
        for col, header in enumerate(headers, 1):
            ws.cell(row=row, column=col, value=header)
        self._apply_header_style(ws, row, 4)
        row += 1
        
        # Collect unique semester-branch combinations
        combos = OrderedDict()
        for slot in self.slots:
            key = (slot.semester, slot.branch_or_program)
            if key not in combos:
                combos[key] = 0
            combos[key] += 1
        
        for (semester, branch), count in combos.items():
            sheet_name = get_sheet_name(semester, branch)
            level = f"Semester {semester_to_roman(semester)}" if semester < 9 else "PG / PhD / MDes"
            ws.cell(row=row, column=1, value=sheet_name).hyperlink = f"#'{sheet_name}'!A1"
            ws.cell(row=row, column=1).font = Font(color="0000FF", underline="single")
            ws.cell(row=row, column=2, value=level)
            ws.cell(row=row, column=3, value=branch)
            ws.cell(row=row, column=4, value=count)
            for col in range(1, 5):
                ws.cell(row=row, column=col).border = self.thin_border
                ws.cell(row=row, column=col).alignment = self.wrap_alignment
            row += 1
        
        # Column widths
        ws.column_dimensions["A"].width = 30
        ws.column_dimensions["B"].width = 25
        ws.column_dimensions["C"].width = 25
        ws.column_dimensions["D"].width = 20
        
        return ws
    
    def build_all_entries_sheet(self) -> openpyxl.worksheet.worksheet.Worksheet:
        """Build the All Entries sheet (normalized flat table)."""
        ws = self.workbook.create_sheet(title="All Entries")
        
        headers = ["Semester", "Branch", "Day", "Course Code", "Instructor", "Room", "Time", "Remarks"]
        for col, header in enumerate(headers, 1):
            ws.cell(row=1, column=col, value=header)
        self._apply_header_style(ws, 1, len(headers))
        
        row = 2
        for slot in sorted(self.slots, key=lambda s: (s.semester, s.branch_or_program, get_day_order_index(s.day), s.start_time or time(0,0))):
            ws.cell(row=row, column=1, value=semester_to_roman(slot.semester) if slot.semester < 9 else "PG")
            ws.cell(row=row, column=2, value=slot.branch_or_program)
            ws.cell(row=row, column=3, value=slot.day)
            ws.cell(row=row, column=4, value=slot.course_code)
            ws.cell(row=row, column=5, value=slot.instructor or "")
            ws.cell(row=row, column=6, value=slot.room or "")
            ws.cell(row=row, column=7, value=time_to_str(slot.start_time))
            ws.cell(row=row, column=8, value=slot.remarks or "")
            
            for col in range(1, 9):
                ws.cell(row=row, column=col).border = self.thin_border
                ws.cell(row=row, column=col).alignment = self.wrap_alignment
            row += 1
        
        # Column widths
        ws.column_dimensions["A"].width = 12
        ws.column_dimensions["B"].width = 25
        ws.column_dimensions["C"].width = 8
        ws.column_dimensions["D"].width = 18
        ws.column_dimensions["E"].width = 18
        ws.column_dimensions["F"].width = 15
        ws.column_dimensions["G"].width = 15
        ws.column_dimensions["H"].width = 40
        
        return ws
    
    def build_check_items_sheet(self) -> openpyxl.worksheet.worksheet.Worksheet:
        """Build the Check Items sheet documenting assumptions."""
        ws = self.workbook.create_sheet(title="Check Items")
        
        ws.cell(row=1, column=1, value="Items to verify against the original timetable").font = self.title_font
        row = 3
        
        headers = ["Topic", "Note"]
        for col, header in enumerate(headers, 1):
            ws.cell(row=row, column=col, value=header)
        self._apply_header_style(ws, row, 2)
        row += 1
        
        check_items = [
            ("Clipped text repaired", "Several elective (OE) cells in the PDF are clipped/overlapping (e.g. \"ME5D03-MZ\", \"C5M03-AK\", \"E8016-HSN\"). They were completed using the full version of the same entry elsewhere in the timetable. Entries marked \"restored\" in Remarks were filled this way."),
            ("Lab rooms not printed", "Lab entries such as NS1002L, ES1002L, EC2002L (Wed/Fri ECE) show no room in the source. Room is left blank rather than guessed."),
            ("ME5D01-SDP", "Printed without a room everywhere in the source; room left blank."),
            ("EC8004-AV (Wed, VII Sem OE9)", "Room is clipped in the source; left blank."),
            ("Fds5001 (Wed, V Sem ECE, 09:00-11:55)", "Printed as \"Fds5001\" in the source; looks like a typo. Kept as printed - please confirm."),
            ("IT3E01-KD (V Sem ECE)", "Printed with \"G1\"/\"G2\" (Mon/Fri). Treated as group G1/G2, not a room."),
            ("SM3010L Batch A (Mon, V Sem)", "Printed as \"CC2-GF\"; shown as CC-GF with a remark."),
            ("Orphan text ignored", "Stray fragments that belong to clipped neighbouring cells were ignored: \"CR101\" (Wed VII ECE/ME 10:00), \"VF-\" (Fri VII ME/SM 09:00), \"M\" (Fri VII DS 09:00), \"CR102\" (Fri VII DS 11:00)."),
            ("Elective (OE) blocks", "OE cells are merged across several branches. Every option listed in the slot is repeated for each branch the cell spans; students choose one."),
            ("Room codes", "L### = lecture hall, CR### = classroom, CC-GF/FF/2F/3F = Computer Centre ground/first/second/third floor, DS (lab) = Design Studio."),
            ("PG / PhD / MDes", "These rows sit under the \"VII Sem\" block of the source, so they are listed under PG."),
            ("Sem I group codes", "CSE-A, CSE-B, ECE (C), ME (D1), SM (D2), DS-E1, DS-E2 as printed in the source."),
        ]
        
        for topic, note in check_items:
            ws.cell(row=row, column=1, value=topic)
            ws.cell(row=row, column=2, value=note)
            for col in range(1, 3):
                ws.cell(row=row, column=col).border = self.thin_border
                ws.cell(row=row, column=col).alignment = self.wrap_alignment
            row += 1
        
        ws.column_dimensions["A"].width = 35
        ws.column_dimensions["B"].width = 100
        
        return ws
    
    def export(self, output_path: Path) -> Dict[str, Any]:
        """Export workbook to file."""
        logger.info(f"Building Excel workbook with {len(self.slots)} slots...")
        
        # Build sheets in order
        self.build_index_sheet()
        
        # Collect unique semester-branch combos in order
        combos = OrderedDict()
        for slot in self.slots:
            key = (slot.semester, slot.branch_or_program)
            if key not in combos:
                combos[key] = True
        
        for (semester, branch) in combos.keys():
            self.build_semester_branch_sheet(semester, branch)
        
        self.build_all_entries_sheet()
        self.build_check_items_sheet()
        
        # Remove default empty sheet if exists
        if "Sheet" in self.workbook.sheetnames:
            self.workbook.remove(self.workbook["Sheet"])
        
        # Save
        self.workbook.save(output_path)
        logger.info(f"Excel saved to {output_path}")
        
        return {
            "path": str(output_path),
            "sheets": self.workbook.sheetnames,
            "total_slots": len(self.slots),
        }


def export_timetable_to_excel(slots: List[TimetableSlot], output_path: Path) -> Dict[str, Any]:
    """Convenience function to export timetable to Excel."""
    exporter = TimetableExcelExporter(slots)
    return exporter.export(output_path)